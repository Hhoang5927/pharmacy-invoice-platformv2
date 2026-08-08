"""
Unit tests for composition_root.review_excel -- pure row<->workbook
mapping, no ServiceContainer/repository/filesystem involved (those are
covered by tests/integration/composition_root/test_cli_export_review.py
and test_cli_import_review.py instead). Deliberately does NOT re-test
SubmitInvoiceReviewUseCase's own parsing/business rules -- that is
already covered by
tests/unit/application/use_cases/test_submit_invoice_review_use_case.py.
"""

from __future__ import annotations

from datetime import date
from decimal import Decimal

import pytest
from openpyxl import load_workbook

from pharmacy_invoice_automation.composition_root import review_excel
from pharmacy_invoice_automation.composition_root.review_excel import (
    ReviewRow,
    build_export_rows,
    read_workbook,
    write_workbook,
)
from pharmacy_invoice_automation.domain.entities.purchase_invoice import PurchaseInvoice
from pharmacy_invoice_automation.domain.entities.purchase_item import PurchaseItem
from pharmacy_invoice_automation.domain.enums.tax_type import TaxType
from pharmacy_invoice_automation.domain.value_objects.money import Money
from pharmacy_invoice_automation.domain.value_objects.quantity import Quantity
from pharmacy_invoice_automation.domain.value_objects.unit import Unit

pytestmark = pytest.mark.unit


def _item(
    item_id: str,
    medicine_name: str = "Paracetamol 500mg",
    *,
    unit_code: str = "hop",
    tax_type: TaxType | None = TaxType.REDUCED,
    retail_units_per_purchase_unit: int | None = None,
    confirmed_website_unit_ratio: Decimal | None = None,
) -> PurchaseItem:
    return PurchaseItem(
        id=item_id,
        medicine_name=medicine_name,
        unit=Unit(code=unit_code),
        quantity=Quantity(Decimal("5")),
        unit_price=Money(Decimal("10000")),
        tax_type=tax_type,
        retail_units_per_purchase_unit=retail_units_per_purchase_unit,
        confirmed_website_unit_ratio=confirmed_website_unit_ratio,
    )


def _invoice(invoice_id: str, invoice_number: str, items: list[PurchaseItem]) -> PurchaseInvoice:
    return PurchaseInvoice(
        id=invoice_id,
        project_id="proj-1",
        invoice_number=invoice_number,
        invoice_date=date(2026, 8, 8),
        items=items,
    )


class TestBuildExportRows:
    def test_one_row_per_purchase_item_across_multiple_invoices(self) -> None:
        invoice_1 = _invoice("inv-1", "INV-001", [_item("item-1"), _item("item-2")])
        invoice_2 = _invoice("inv-2", "INV-002", [_item("item-3")])

        rows = build_export_rows(
            [invoice_1, invoice_2],
            issues_by_invoice_id={},
            supplier_name_by_invoice_id={},
        )

        assert [r.item_id for r in rows] == ["item-1", "item-2", "item-3"]
        assert [r.invoice_id for r in rows] == ["inv-1", "inv-1", "inv-2"]

    def test_issues_are_joined_and_repeated_on_every_row_of_that_invoice(self) -> None:
        invoice = _invoice("inv-1", "INV-001", [_item("item-1"), _item("item-2")])

        rows = build_export_rows(
            [invoice],
            issues_by_invoice_id={"inv-1": ["Vấn đề A", "Vấn đề B"]},
            supplier_name_by_invoice_id={},
        )

        assert rows[0].issues == "Vấn đề A; Vấn đề B"
        assert rows[1].issues == "Vấn đề A; Vấn đề B"

    def test_supplier_name_resolved_per_invoice(self) -> None:
        invoice = _invoice("inv-1", "INV-001", [_item("item-1")])

        rows = build_export_rows(
            [invoice],
            issues_by_invoice_id={},
            supplier_name_by_invoice_id={"inv-1": "Công ty Dược ABC"},
        )

        assert rows[0].supplier_name == "Công ty Dược ABC"

    @pytest.mark.parametrize(
        ("tax_type", "expected_label"),
        [
            (TaxType.STANDARD, "10%"),
            (TaxType.REDUCED, "5%"),
            (TaxType.EXEMPT, "0%"),
            (TaxType.EIGHT_PERCENT, "8%"),
            (TaxType.OTHER, "Khác"),
            (None, ""),
        ],
    )
    def test_vat_label_covers_every_tax_type_and_none(
        self, tax_type: TaxType | None, expected_label: str
    ) -> None:
        invoice = _invoice("inv-1", "INV-001", [_item("item-1", tax_type=tax_type)])

        rows = build_export_rows([invoice], issues_by_invoice_id={}, supplier_name_by_invoice_id={})

        assert rows[0].vat_label == expected_label

    def test_current_conversion_fields_blank_when_none(self) -> None:
        invoice = _invoice("inv-1", "INV-001", [_item("item-1")])

        rows = build_export_rows([invoice], issues_by_invoice_id={}, supplier_name_by_invoice_id={})

        assert rows[0].current_retail_ratio == ""
        assert rows[0].current_website_ratio == ""

    def test_current_conversion_fields_populated_when_set(self) -> None:
        invoice = _invoice(
            "inv-1",
            "INV-001",
            [
                _item(
                    "item-1",
                    retail_units_per_purchase_unit=20,
                    confirmed_website_unit_ratio=Decimal("1"),
                )
            ],
        )

        rows = build_export_rows([invoice], issues_by_invoice_id={}, supplier_name_by_invoice_id={})

        assert rows[0].current_retail_ratio == "20"
        assert rows[0].current_website_ratio == "1"


class TestWriteWorkbook:
    def test_header_row_matches_documented_column_order_and_labels(self) -> None:
        workbook_bytes = write_workbook([])
        workbook = load_workbook(__import__("io").BytesIO(workbook_bytes))
        sheet = workbook.active

        header = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]

        assert header == [label for _, label in review_excel._COLUMNS]

    def test_round_trips_row_values_through_openpyxl(self) -> None:
        row = ReviewRow(
            invoice_id="inv-1",
            item_id="item-1",
            invoice_number="INV-001",
            invoice_date="2026-08-08",
            supplier_name="Công ty Dược ABC",
            medicine_name="Paracetamol 500mg",
            unit_code="hop",
            quantity="5",
            unit_price="10,000 VND",
            vat_label="5%",
            current_retail_ratio="",
            current_website_ratio="",
            issues="Vấn đề A",
        )

        workbook = load_workbook(__import__("io").BytesIO(write_workbook([row])))
        sheet = workbook.active
        data_row = [cell.value for cell in next(sheet.iter_rows(min_row=2, max_row=2))]

        assert data_row[review_excel._COLUMN_INDEX["invoice_id"]] == "inv-1"
        assert data_row[review_excel._COLUMN_INDEX["item_id"]] == "item-1"
        assert data_row[review_excel._COLUMN_INDEX["medicine_name"]] == "Paracetamol 500mg"
        assert data_row[review_excel._COLUMN_INDEX["issues"]] == "Vấn đề A"
        # Correction/decision columns are always blank in a fresh export.
        assert data_row[review_excel._COLUMN_INDEX["correction_retail_ratio"]] is None
        assert data_row[review_excel._COLUMN_INDEX["decision"]] is None

    def test_adds_dropdown_validation_only_for_decision_and_medicine_type(self) -> None:
        row = ReviewRow(
            invoice_id="inv-1",
            item_id="item-1",
            invoice_number="INV-001",
            invoice_date="2026-08-08",
            supplier_name="",
            medicine_name="Paracetamol 500mg",
            unit_code="hop",
            quantity="5",
            unit_price="10,000 VND",
            vat_label="5%",
            current_retail_ratio="",
            current_website_ratio="",
            issues="",
        )

        workbook = load_workbook(__import__("io").BytesIO(write_workbook([row])))
        sheet = workbook.active

        validated_columns = set()
        for dv in sheet.data_validations.dataValidation:
            for cell_range in dv.sqref.ranges:
                validated_columns.add(cell_range.min_col)

        decision_col = review_excel._COLUMN_INDEX["decision"] + 1
        medicine_type_col = review_excel._COLUMN_INDEX["correction_medicine_type"] + 1
        retail_unit_override_col = (
            review_excel._COLUMN_INDEX["correction_retail_unit_override"] + 1
        )

        assert decision_col in validated_columns
        assert medicine_type_col in validated_columns
        assert retail_unit_override_col not in validated_columns


def _make_export_bytes(rows: list[ReviewRow]) -> bytes:
    return write_workbook(rows)


def _load_and_edit(
    xlsx_bytes: bytes, edits: dict[tuple[int, str], str]
) -> bytes:
    """edits: {(row_number_1_indexed_from_2, column_key): value}"""
    import io

    workbook = load_workbook(io.BytesIO(xlsx_bytes))
    sheet = workbook.active
    for (row_number, column_key), value in edits.items():
        sheet.cell(row=row_number, column=review_excel._COLUMN_INDEX[column_key] + 1, value=value)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _base_rows() -> list[ReviewRow]:
    return [
        ReviewRow(
            invoice_id="inv-1",
            item_id="item-1",
            invoice_number="INV-001",
            invoice_date="2026-08-08",
            supplier_name="Công ty Dược ABC",
            medicine_name="Paracetamol 500mg",
            unit_code="hop",
            quantity="5",
            unit_price="10,000 VND",
            vat_label="5%",
            current_retail_ratio="",
            current_website_ratio="",
            issues="",
        ),
        ReviewRow(
            invoice_id="inv-1",
            item_id="item-2",
            invoice_number="INV-001",
            invoice_date="2026-08-08",
            supplier_name="Công ty Dược ABC",
            medicine_name="Amoxicillin 500mg",
            unit_code="hop",
            quantity="2",
            unit_price="20,000 VND",
            vat_label="5%",
            current_retail_ratio="",
            current_website_ratio="",
            issues="",
        ),
    ]


class TestReadWorkbookSkippingAndDecision:
    def test_blank_decision_on_every_row_yields_skipped_not_error(self) -> None:
        result = read_workbook(_make_export_bytes(_base_rows()))

        assert result.reviews == []
        assert result.errors == []
        assert result.skipped_invoice_ids == ["inv-1"]

    def test_duyet_maps_to_reviewer_approved_true(self) -> None:
        edited = _load_and_edit(
            _make_export_bytes(_base_rows()), {(2, "decision"): "duyet"}
        )

        result = read_workbook(edited)

        assert len(result.reviews) == 1
        assert result.reviews[0].reviewer_approved is True
        assert result.reviews[0].invoice_id == "inv-1"

    def test_tu_choi_maps_to_reviewer_approved_false(self) -> None:
        edited = _load_and_edit(
            _make_export_bytes(_base_rows()), {(2, "decision"): "tu_choi"}
        )

        result = read_workbook(edited)

        assert len(result.reviews) == 1
        assert result.reviews[0].reviewer_approved is False

    def test_invalid_decision_text_is_reported_as_an_error_not_silently_skipped(
        self,
    ) -> None:
        edited = _load_and_edit(
            _make_export_bytes(_base_rows()), {(2, "decision"): "OK"}
        )

        result = read_workbook(edited)

        assert result.reviews == []
        assert len(result.errors) == 1
        assert result.errors[0].field_label == "Quyết định"
        assert result.errors[0].raw_value == "OK"

    def test_conflicting_decision_across_rows_of_same_invoice_is_an_error(self) -> None:
        edited = _load_and_edit(
            _make_export_bytes(_base_rows()),
            {(2, "decision"): "duyet", (3, "decision"): "tu_choi"},
        )

        result = read_workbook(edited)

        assert result.reviews == []
        assert len(result.errors) == 1
        assert result.errors[0].field_label == "Quyết định"


class TestReadWorkbookItemLevelFields:
    def test_builds_item_level_corrected_field_keys_matching_use_case_vocabulary(
        self,
    ) -> None:
        edited = _load_and_edit(
            _make_export_bytes(_base_rows()),
            {
                (2, "decision"): "duyet",
                (2, "correction_retail_ratio"): "10",
                (3, "correction_medicine_type"): "prescription",
            },
        )

        result = read_workbook(edited)

        assert len(result.reviews) == 1
        fields = result.reviews[0].corrected_fields
        assert fields["item.item-1.retail_units_per_purchase_unit"] == "10"
        assert fields["item.item-2.medicine_type"] == "prescription"

    def test_invalid_item_level_value_excludes_the_whole_invoice_and_is_reported(
        self,
    ) -> None:
        edited = _load_and_edit(
            _make_export_bytes(_base_rows()),
            {
                (2, "decision"): "duyet",
                (2, "correction_retail_ratio"): "not-a-number",
                (3, "correction_medicine_type"): "prescription",
            },
        )

        result = read_workbook(edited)

        assert result.reviews == []
        assert len(result.errors) == 1
        assert result.errors[0].medicine_name == "Paracetamol 500mg"
        assert result.errors[0].field_label == "SỬA: Hệ số quy đổi Viên"

    def test_invalid_retail_unit_override_code_is_reported(self) -> None:
        edited = _load_and_edit(
            _make_export_bytes(_base_rows()),
            {
                (2, "decision"): "duyet",
                (2, "correction_retail_unit_override"): "not_a_real_code",
            },
        )

        result = read_workbook(edited)

        assert result.reviews == []
        assert result.errors[0].field_label == "SỬA: Mã ĐVT bán lẻ ghi đè"

    def test_valid_retail_unit_override_code_is_accepted(self) -> None:
        edited = _load_and_edit(
            _make_export_bytes(_base_rows()),
            {(2, "decision"): "duyet", (2, "correction_retail_unit_override"): "lo"},
        )

        result = read_workbook(edited)

        assert len(result.reviews) == 1
        assert result.reviews[0].corrected_fields["item.item-1.retail_unit_override"] == "lo"


class TestReadWorkbookInvoiceLevelFields:
    def test_first_row_only_filled_is_applied(self) -> None:
        edited = _load_and_edit(
            _make_export_bytes(_base_rows()),
            {(2, "decision"): "duyet", (2, "correction_invoice_number"): "INV-001-FIXED"},
        )

        result = read_workbook(edited)

        assert len(result.reviews) == 1
        assert result.reviews[0].corrected_fields["invoice_number"] == "INV-001-FIXED"

    def test_conflicting_invoice_level_value_across_rows_is_an_error(self) -> None:
        edited = _load_and_edit(
            _make_export_bytes(_base_rows()),
            {
                (2, "decision"): "duyet",
                (2, "correction_invoice_number"): "INV-001-A",
                (3, "correction_invoice_number"): "INV-001-B",
            },
        )

        result = read_workbook(edited)

        assert result.reviews == []
        assert any(e.field_label == "SỬA: Số hóa đơn" for e in result.errors)

    def test_invalid_invoice_date_format_is_reported(self) -> None:
        edited = _load_and_edit(
            _make_export_bytes(_base_rows()),
            {(2, "decision"): "duyet", (2, "correction_invoice_date"): "08/08/2026"},
        )

        result = read_workbook(edited)

        assert result.reviews == []
        assert any(e.field_label == "SỬA: Ngày hóa đơn (YYYY-MM-DD)" for e in result.errors)

    def test_excel_native_date_cell_is_normalized_to_iso_string(self) -> None:
        edited = _load_and_edit(
            _make_export_bytes(_base_rows()),
            {(2, "decision"): "duyet", (2, "correction_invoice_date"): "2026-09-01"},
        )
        # Simulate Excel storing this as a real date cell type by writing
        # a datetime.date object directly, as a spreadsheet app would after
        # the PO reformats the column as a date.
        import io

        workbook = load_workbook(io.BytesIO(edited))
        sheet = workbook.active
        sheet.cell(
            row=2,
            column=review_excel._COLUMN_INDEX["correction_invoice_date"] + 1,
            value=date(2026, 9, 1),
        )
        buffer = io.BytesIO()
        workbook.save(buffer)

        result = read_workbook(buffer.getvalue())

        assert len(result.reviews) == 1
        assert result.reviews[0].corrected_fields["invoice_date"] == "2026-09-01"


class TestReadWorkbookRoundTrip:
    def test_export_then_import_with_no_edits_yields_no_reviews_and_a_skip(self) -> None:
        invoice = _invoice("inv-1", "INV-001", [_item("item-1")])
        rows = build_export_rows([invoice], issues_by_invoice_id={}, supplier_name_by_invoice_id={})

        result = read_workbook(write_workbook(rows))

        assert result.reviews == []
        assert result.errors == []
        assert result.skipped_invoice_ids == ["inv-1"]

    def test_two_different_invoices_are_grouped_and_processed_independently(self) -> None:
        invoice_1 = _invoice("inv-1", "INV-001", [_item("item-1")])
        invoice_2 = _invoice("inv-2", "INV-002", [_item("item-2")])
        rows = build_export_rows(
            [invoice_1, invoice_2], issues_by_invoice_id={}, supplier_name_by_invoice_id={}
        )
        xlsx = write_workbook(rows)
        edited = _load_and_edit(xlsx, {(2, "decision"): "duyet", (3, "decision"): "tu_choi"})

        result = read_workbook(edited)

        by_id = {r.invoice_id: r for r in result.reviews}
        assert by_id["inv-1"].reviewer_approved is True
        assert by_id["inv-2"].reviewer_approved is False

    def test_duplicate_invoice_number_across_two_different_invoices_does_not_merge_them(
        self,
    ) -> None:
        """Two genuinely different invoices sharing the same (duplicate,
        under-review) invoice number must still be grouped/applied
        independently -- grouping is by the hidden invoice_id, never the
        human-readable invoice_number text column."""
        invoice_1 = _invoice("inv-1", "DUP-001", [_item("item-1")])
        invoice_2 = _invoice("inv-2", "DUP-001", [_item("item-2")])
        rows = build_export_rows(
            [invoice_1, invoice_2], issues_by_invoice_id={}, supplier_name_by_invoice_id={}
        )
        xlsx = write_workbook(rows)
        edited = _load_and_edit(xlsx, {(2, "decision"): "duyet", (3, "decision"): "tu_choi"})

        result = read_workbook(edited)

        by_id = {r.invoice_id: r for r in result.reviews}
        assert by_id["inv-1"].reviewer_approved is True
        assert by_id["inv-2"].reviewer_approved is False
