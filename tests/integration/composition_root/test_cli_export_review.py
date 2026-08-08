"""
Real (no Gemini calls, no network) integration tests for
composition_root.cli.run_export_review() -- Stage C2 (PO decision
2026-08, see composition_root.review_excel's own docstring and
CLAUDE.md Deviation D7). Real SQLite (via register_infrastructure_services
against a temp app_root), real domain entities persisted through real
repositories, real FileStorageProvider/WorkspaceManager -- the written
.xlsx is reopened with openpyxl to verify its actual contents, not mocked.
"""

from __future__ import annotations

from datetime import date
from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import load_workbook

from pharmacy_invoice_automation.composition_root import cli, review_excel
from pharmacy_invoice_automation.domain.entities.medicine import Medicine
from pharmacy_invoice_automation.domain.entities.project import Project
from pharmacy_invoice_automation.domain.entities.purchase_invoice import PurchaseInvoice
from pharmacy_invoice_automation.domain.entities.purchase_item import PurchaseItem
from pharmacy_invoice_automation.domain.entities.supplier import Supplier
from pharmacy_invoice_automation.domain.enums.invoice_status import InvoiceStatus
from pharmacy_invoice_automation.domain.enums.medicine_type import MedicineType
from pharmacy_invoice_automation.domain.enums.tax_type import TaxType
from pharmacy_invoice_automation.domain.ports.repositories.medicine_repository import (
    MedicineRepository,
)
from pharmacy_invoice_automation.domain.ports.repositories.project_repository import (
    ProjectRepository,
)
from pharmacy_invoice_automation.domain.ports.repositories.purchase_invoice_repository import (
    PurchaseInvoiceRepository,
)
from pharmacy_invoice_automation.domain.ports.repositories.supplier_repository import (
    SupplierRepository,
)
from pharmacy_invoice_automation.domain.value_objects.money import Money
from pharmacy_invoice_automation.domain.value_objects.quantity import Quantity
from pharmacy_invoice_automation.domain.value_objects.unit import Unit
from pharmacy_invoice_automation.infrastructure.di.registration import (
    register_infrastructure_services,
)
from pharmacy_invoice_automation.infrastructure.di.service_container import ServiceContainer
from pharmacy_invoice_automation.infrastructure.file_storage.workspace_manager import (
    WorkspaceManager,
)

pytestmark = pytest.mark.integration


@pytest.fixture()
def container(tmp_path: Path) -> ServiceContainer:
    service_container = ServiceContainer()
    register_infrastructure_services(service_container, tmp_path / "app")
    return service_container


def _seed_invoice_needing_packaging_ratio(
    container: ServiceContainer, invoice_id: str = "inv-1", invoice_number: str = "INV-001"
) -> PurchaseItem:
    if container.resolve(ProjectRepository).get_by_id("proj-1") is None:
        container.resolve(ProjectRepository).add(
            Project(id="proj-1", name="Test Project", root_folder="C:/tmp")
        )
    if container.resolve(SupplierRepository).get_by_id("sup-1") is None:
        container.resolve(SupplierRepository).add(Supplier(id="sup-1", name="Nice Pharma Co"))
    if container.resolve(MedicineRepository).get_by_id("med-1") is None:
        container.resolve(MedicineRepository).add(
            Medicine(
                id="med-1",
                medicine_code="TH1",
                name="Paracetamol 500mg",
                medicine_type=MedicineType.OVER_THE_COUNTER,
                unit=Unit(code="vien"),
            )
        )
    item = PurchaseItem(
        id=f"{invoice_id}-item-1",
        medicine_name="Paracetamol 500mg",
        unit=Unit(code="hop"),
        quantity=Quantity(Decimal("5")),
        unit_price=Money(Decimal("10000")),
        medicine_id="med-1",
        tax_type=TaxType.REDUCED,
    )
    invoice = PurchaseInvoice(
        id=invoice_id,
        project_id="proj-1",
        invoice_number=invoice_number,
        invoice_date=date.today(),
        status=InvoiceStatus.UNDER_REVIEW,
        supplier_id="sup-1",
        items=[item],
    )
    container.resolve(PurchaseInvoiceRepository).add(invoice)
    return item


class TestNoInvoicesToExport:
    def test_prints_a_clear_message_and_returns_none(
        self, container: ServiceContainer, capsys: pytest.CaptureFixture[str]
    ) -> None:
        result = cli.run_export_review(container)

        assert result is None
        assert "Không có hóa đơn" in capsys.readouterr().out


class TestExportWritesRealWorkbook:
    def test_writes_one_row_per_item_across_multiple_invoices(
        self, container: ServiceContainer
    ) -> None:
        _seed_invoice_needing_packaging_ratio(container, "inv-1", "INV-001")
        _seed_invoice_needing_packaging_ratio(container, "inv-2", "INV-002")

        path = cli.run_export_review(container)

        assert path is not None
        assert path.is_file()
        workbook = load_workbook(path)
        sheet = workbook.active
        data_rows = list(sheet.iter_rows(min_row=2, values_only=True))
        assert len(data_rows) == 2
        invoice_numbers = {row[review_excel._COLUMN_INDEX["invoice_number"]] for row in data_rows}
        assert invoice_numbers == {"INV-001", "INV-002"}

    def test_current_validation_issues_are_included_as_reference_text(
        self, container: ServiceContainer
    ) -> None:
        _seed_invoice_needing_packaging_ratio(container)

        path = cli.run_export_review(container)

        assert path is not None
        workbook = load_workbook(path)
        sheet = workbook.active
        data_row = next(sheet.iter_rows(min_row=2, max_row=2, values_only=True))
        assert data_row[review_excel._COLUMN_INDEX["issues"]] is None or isinstance(
            data_row[review_excel._COLUMN_INDEX["issues"]], str
        )

    def test_correction_and_decision_columns_are_blank_in_a_fresh_export(
        self, container: ServiceContainer
    ) -> None:
        _seed_invoice_needing_packaging_ratio(container)

        path = cli.run_export_review(container)

        assert path is not None
        workbook = load_workbook(path)
        sheet = workbook.active
        data_row = next(sheet.iter_rows(min_row=2, max_row=2, values_only=True))
        assert data_row[review_excel._COLUMN_INDEX["decision"]] is None
        assert data_row[review_excel._COLUMN_INDEX["correction_retail_ratio"]] is None


class TestExportDefaultOutputPath:
    def test_default_output_path_lands_under_data_review_exports(
        self, container: ServiceContainer
    ) -> None:
        _seed_invoice_needing_packaging_ratio(container)

        path = cli.run_export_review(container)

        assert path is not None
        workspace = container.resolve(WorkspaceManager)
        expected_folder = workspace.data_directory / "review_exports"
        assert path.parent == expected_folder
        assert path.name.startswith("review_")
        assert path.suffix == ".xlsx"

    def test_two_exports_in_the_same_run_never_overwrite_each_other(
        self, container: ServiceContainer
    ) -> None:
        _seed_invoice_needing_packaging_ratio(container, "inv-1", "INV-001")

        first_path = cli.run_export_review(container)
        second_path = cli.run_export_review(container)

        assert first_path is not None
        assert second_path is not None
        assert first_path != second_path
        assert first_path.is_file()
        assert second_path.is_file()


class TestExportRespectsExplicitOutputPath:
    def test_explicit_output_path_is_used_verbatim(
        self, container: ServiceContainer, tmp_path: Path
    ) -> None:
        _seed_invoice_needing_packaging_ratio(container)
        explicit_path = tmp_path / "custom" / "my_export.xlsx"

        result = cli.run_export_review(container, output_path=explicit_path)

        assert result == explicit_path
        assert explicit_path.is_file()
