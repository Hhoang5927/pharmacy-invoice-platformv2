"""
Real (no Gemini calls, no network) integration tests for
composition_root.cli.run_import_review() -- Stage C2 (PO decision
2026-08). Real SQLite, real domain entities, real
SubmitInvoiceReviewUseCase/InvoiceValidator -- the whole point of these
tests is proving run_import_review() only ever builds a
SubmitInvoiceReviewCommand and calls the real use case, never
reimplementing any correction/validation logic itself.
"""

from __future__ import annotations

from datetime import date
from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import load_workbook

from pharmacy_invoice_automation.composition_root import cli, review_excel
from pharmacy_invoice_automation.domain.entities.medicine import Medicine
from pharmacy_invoice_automation.domain.entities.project import Project
from pharmacy_invoice_automation.domain.entities.purchase_invoice import PurchaseInvoice
from pharmacy_invoice_automation.domain.entities.purchase_item import PurchaseItem
from pharmacy_invoice_automation.domain.entities.supplier import Supplier
from pharmacy_invoice_automation.domain.enums.invoice_status import InvoiceStatus
from pharmacy_invoice_automation.domain.enums.medicine_type import MedicineType
from pharmacy_invoice_automation.domain.enums.tax_type import TaxType
from pharmacy_invoice_automation.domain.ports.repositories.medicine_repository import (
    MedicineRepository,
)
from pharmacy_invoice_automation.domain.ports.repositories.project_repository import (
    ProjectRepository,
)
from pharmacy_invoice_automation.domain.ports.repositories.purchase_invoice_repository import (
    PurchaseInvoiceRepository,
)
from pharmacy_invoice_automation.domain.ports.repositories.supplier_repository import (
    SupplierRepository,
)
from pharmacy_invoice_automation.domain.value_objects.money import Money
from pharmacy_invoice_automation.domain.value_objects.quantity import Quantity
from pharmacy_invoice_automation.domain.value_objects.unit import Unit
from pharmacy_invoice_automation.infrastructure.di.registration import (
    register_infrastructure_services,
)
from pharmacy_invoice_automation.infrastructure.di.service_container import ServiceContainer

pytestmark = pytest.mark.integration


@pytest.fixture()
def container(tmp_path: Path) -> ServiceContainer:
    service_container = ServiceContainer()
    register_infrastructure_services(service_container, tmp_path / "app")
    return service_container


def _seed_invoice_needing_packaging_ratio(container: ServiceContainer) -> PurchaseItem:
    container.resolve(ProjectRepository).add(
        Project(id="proj-1", name="Test Project", root_folder="C:/tmp")
    )
    container.resolve(SupplierRepository).add(Supplier(id="sup-1", name="Nice Pharma Co"))
    container.resolve(MedicineRepository).add(
        Medicine(
            id="med-1",
            medicine_code="TH1",
            name="Paracetamol 500mg",
            medicine_type=MedicineType.OVER_THE_COUNTER,
            unit=Unit(code="vien"),
        )
    )
    item = PurchaseItem(
        id="item-1",
        medicine_name="Paracetamol 500mg",
        unit=Unit(code="hop"),
        quantity=Quantity(Decimal("5")),
        unit_price=Money(Decimal("10000")),
        medicine_id="med-1",
        tax_type=TaxType.REDUCED,
    )
    invoice = PurchaseInvoice(
        id="inv-1",
        project_id="proj-1",
        invoice_number="INV-001",
        invoice_date=date.today(),
        status=InvoiceStatus.UNDER_REVIEW,
        supplier_id="sup-1",
        items=[item],
    )
    container.resolve(PurchaseInvoiceRepository).add(invoice)
    return item


def _edit_and_save(xlsx_path: Path, edits: dict[tuple[int, str], str]) -> Path:
    workbook = load_workbook(xlsx_path)
    sheet = workbook.active
    for (row_number, column_key), value in edits.items():
        sheet.cell(row=row_number, column=review_excel._COLUMN_INDEX[column_key] + 1, value=value)
    workbook.save(xlsx_path)
    return xlsx_path


class TestExportEditImportApprovesAndAdvances:
    def test_full_round_trip_advances_the_invoice_to_ready_for_import(
        self, container: ServiceContainer
    ) -> None:
        _seed_invoice_needing_packaging_ratio(container)
        exported_path = cli.run_export_review(container)
        assert exported_path is not None

        _edit_and_save(
            exported_path,
            {
                (2, "correction_retail_ratio"): "10",
                (2, "decision"): "duyet",
            },
        )

        results = cli.run_import_review(container, exported_path)

        assert len(results) == 1
        assert results[0].is_success, results[0].errors
        assert results[0].value is not None
        assert results[0].value.status == InvoiceStatus.READY_FOR_IMPORT.value

        reloaded = container.resolve(PurchaseInvoiceRepository).get_by_id("inv-1")
        assert reloaded is not None
        assert reloaded.status is InvoiceStatus.READY_FOR_IMPORT
        assert reloaded.items[0].retail_units_per_purchase_unit == 10

    def test_the_correction_is_learned_onto_the_medicine_for_next_time(
        self, container: ServiceContainer
    ) -> None:
        _seed_invoice_needing_packaging_ratio(container)
        exported_path = cli.run_export_review(container)
        assert exported_path is not None
        _edit_and_save(
            exported_path,
            {(2, "correction_retail_ratio"): "10", (2, "decision"): "duyet"},
        )

        cli.run_import_review(container, exported_path)

        medicine = container.resolve(MedicineRepository).get_by_id("med-1")
        assert medicine is not None
        assert medicine.retail_units_per_purchase_unit == 10


class TestRejectingViaExcel:
    def test_tu_choi_leaves_the_invoice_under_review(self, container: ServiceContainer) -> None:
        _seed_invoice_needing_packaging_ratio(container)
        exported_path = cli.run_export_review(container)
        assert exported_path is not None
        _edit_and_save(exported_path, {(2, "decision"): "tu_choi"})

        results = cli.run_import_review(container, exported_path)

        assert len(results) == 1
        assert results[0].is_success
        reloaded = container.resolve(PurchaseInvoiceRepository).get_by_id("inv-1")
        assert reloaded is not None
        assert reloaded.status is InvoiceStatus.UNDER_REVIEW


class TestBlankDecisionSkipsWithoutCallingUseCase:
    def test_blank_decision_leaves_invoice_completely_untouched(
        self, container: ServiceContainer
    ) -> None:
        _seed_invoice_needing_packaging_ratio(container)
        exported_path = cli.run_export_review(container)
        assert exported_path is not None
        before = container.resolve(PurchaseInvoiceRepository).get_by_id("inv-1")
        assert before is not None
        before_version = before.version

        results = cli.run_import_review(container, exported_path)

        assert results == []
        after = container.resolve(PurchaseInvoiceRepository).get_by_id("inv-1")
        assert after is not None
        assert after.status is InvoiceStatus.UNDER_REVIEW
        assert after.version == before_version


class TestInvalidCellReportedAndInvoiceNotApplied:
    def test_invalid_value_is_not_silently_ignored_and_nothing_is_applied(
        self, container: ServiceContainer, capsys: pytest.CaptureFixture[str]
    ) -> None:
        _seed_invoice_needing_packaging_ratio(container)
        exported_path = cli.run_export_review(container)
        assert exported_path is not None
        _edit_and_save(
            exported_path,
            {
                (2, "decision"): "duyet",
                (2, "correction_retail_ratio"): "not-a-number",
            },
        )
        before = container.resolve(PurchaseInvoiceRepository).get_by_id("inv-1")
        assert before is not None
        before_version = before.version

        results = cli.run_import_review(container, exported_path)

        assert results == []
        output = capsys.readouterr().out
        assert "Các lỗi cần sửa lại" in output
        assert "Paracetamol 500mg" in output
        after = container.resolve(PurchaseInvoiceRepository).get_by_id("inv-1")
        assert after is not None
        assert after.status is InvoiceStatus.UNDER_REVIEW
        assert after.version == before_version


class TestStaleInvoiceAlreadyResolved:
    def test_invoice_no_longer_under_review_reports_a_clean_error(
        self, container: ServiceContainer
    ) -> None:
        _seed_invoice_needing_packaging_ratio(container)
        exported_path = cli.run_export_review(container)
        assert exported_path is not None
        _edit_and_save(exported_path, {(2, "decision"): "duyet"})

        repository = container.resolve(PurchaseInvoiceRepository)
        invoice = repository.get_by_id("inv-1")
        assert invoice is not None
        invoice.transition_to(InvoiceStatus.READY_FOR_IMPORT)
        repository.update(invoice)

        results = cli.run_import_review(container, exported_path)

        assert len(results) == 1
        assert not results[0].is_success
        assert any("not awaiting review" in e for e in results[0].errors)
